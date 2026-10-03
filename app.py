import os
import tempfile
import pandas as pd
import streamlit as st
from dotenv import load_dotenv

from pypdf import PdfReader
from docx import Document

from langchain.schema import Document as LangDocument
from langchain.memory import ConversationBufferMemory
from langchain.chains import ConversationalRetrievalChain
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_pinecone import PineconeVectorStore
from pinecone import Pinecone, ServerlessSpec
from langchain_groq import ChatGroq

# =========================
# PAGE CONFIGURATION
# =========================
st.set_page_config(
    page_title="AI Support",
    page_icon="✨",
    layout="wide",
    initial_sidebar_state="expanded" 
)

# =========================
# HELPER: HEX TO RGB FOR GLOWS
# =========================
def hex_to_rgb(hex_color):
    """Converts a hex color string to an RGB tuple so we can use rgba() in CSS."""
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

# =========================
# SIDEBAR: SETTINGS & THEMING
# =========================
with st.sidebar:
    st.title("⚙️ Workspace Settings")
    
    # --- DYNAMIC THEMING CONTROLS ---
    st.markdown("### 🎨 Theme Customization")
    accent_color = st.color_picker("Primary Accent Color", "#007BFF") # Default Blue
    bg_gradient_1 = st.color_picker("Background Gradient (Top)", "#F8F9FA") # Default Light
    bg_gradient_2 = st.color_picker("Background Gradient (Bottom)", "#E9ECEF") # Default Light
    
    # Calculate RGBA values for glassmorphism based on chosen accent
    r, g, b = hex_to_rgb(accent_color)
    accent_glow = f"rgba({r}, {g}, {b}, 0.3)"
    accent_bg_light = f"rgba({r}, {g}, {b}, 0.05)"
    accent_border = f"rgba({r}, {g}, {b}, 0.2)"
    accent_button_hover = f"rgba({r}, {g}, {b}, 0.1)"
    
    text_color = "#1E1E1E" if int(bg_gradient_1.lstrip('#')[:2], 16) > 128 else "#F8FAFC"
    
    st.markdown("<hr>", unsafe_allow_html=True)
    
    st.markdown("### 👤 User Profile")
    active_user = st.selectbox(
        "Active Session", 
        ["User 1 (Alice)", "User 2 (Bob)", "User 3 (Charlie)"],
        label_visibility="collapsed"
    )
    
    st.markdown("<hr>", unsafe_allow_html=True)
    
    st.markdown("### 📁 Knowledge Base")
    uploaded_files = st.file_uploader(
        "Inject context documents into the RAG pipeline",
        type=["pdf", "docx", "csv", "txt"],
        accept_multiple_files=True
    )

# =========================
# DYNAMIC GLASSMORPHISM CSS
# =========================
st.markdown(f"""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600&display=swap');

    /* Background and Base Text */
    .stApp {{
        background: linear-gradient(135deg, {bg_gradient_1} 0%, {bg_gradient_2} 100%);
        background-attachment: fixed;
        color: {text_color} !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
    }}
    
    p, span, div, label {{
        color: {text_color} !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
    }}

    /* Sidebar Glass Effect */
    [data-testid="stSidebar"] {{
        background-color: rgba(255, 255, 255, 0.1) !important;
        backdrop-filter: blur(15px) !important;
        -webkit-backdrop-filter: blur(15px) !important;
        border-right: 1px solid rgba(128, 128, 128, 0.2);
    }}

    /* Accent Headers */
    h1, h2, h3 {{
        color: {accent_color} !important;
        font-weight: 600 !important;
        letter-spacing: -0.5px;
        text-shadow: 0px 0px 15px {accent_glow};
    }}

    /* Chat Messages Glass Bubbles */
    .stChatMessage:has([data-testid="chatAvatarIcon-user"]) {{
        background: rgba(128, 128, 128, 0.05) !important;
        backdrop-filter: blur(10px) !important;
        border: 1px solid rgba(128, 128, 128, 0.1) !important;
        border-radius: 16px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05);
        margin-bottom: 20px;
        padding: 15px 20px;
    }}
    
    .stChatMessage:has([data-testid="chatAvatarIcon-assistant"]) {{
        background: {accent_bg_light} !important;
        backdrop-filter: blur(12px) !important;
        border: 1px solid {accent_border} !important;
        border-radius: 16px;
        box-shadow: 0 10px 25px -5px {accent_glow};
        margin-bottom: 20px;
        padding: 15px 20px;
    }}

    /* Buttons */
    .stButton>button {{
        background: {accent_bg_light} !important;
        color: {accent_color} !important;
        border: 1px solid {accent_border} !important;
        border-radius: 12px !important;
        font-weight: 500 !important;
        backdrop-filter: blur(5px);
        transition: all 0.2s ease-in-out !important;
    }}
    .stButton>button:hover {{
        background: {accent_button_hover} !important;
        border: 1px solid {accent_color} !important;
        transform: translateY(-2px);
        box-shadow: 0 0 15px {accent_glow};
    }}

    /* Input Chat Bar */
    .stChatInputContainer {{
        padding-bottom: 20px;
    }}
    .stChatInputContainer textarea {{
        background-color: rgba(128, 128, 128, 0.1) !important;
        backdrop-filter: blur(10px);
        color: {text_color} !important;
        border-radius: 16px !important;
        border: 1px solid rgba(128, 128, 128, 0.2) !important;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05);
    }}
    .stChatInputContainer textarea:focus {{
        border-color: {accent_color} !important;
        box-shadow: 0 0 15px {accent_glow} !important;
    }}
    </style>
""", unsafe_allow_html=True)

# =========================
# LOAD ENV VARIABLES
# =========================
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
PINECONE_INDEX = os.getenv("PINECONE_INDEX", "customer-support-rag")

if not GROQ_API_KEY or not PINECONE_API_KEY:
    st.error("🚨 Missing API Keys! Please check your .env file.")
    st.stop()

# =========================
# MULTI-USER SESSION STATE
# =========================
if "user_profiles" not in st.session_state:
    st.session_state.user_profiles = {}

if active_user not in st.session_state.user_profiles:
    st.session_state.user_profiles[active_user] = {
        "chat_history": [],
        "memory": ConversationBufferMemory(memory_key="chat_history", return_messages=True)
    }

user_data = st.session_state.user_profiles[active_user]

# =========================
# CACHED INITIALIZATION
# =========================
@st.cache_resource(show_spinner=False)
def init_backend():
    pc = Pinecone(api_key=PINECONE_API_KEY)
    existing_indexes = [index["name"] for index in pc.list_indexes()]
    if PINECONE_INDEX not in existing_indexes:
        pc.create_index(
            name=PINECONE_INDEX,
            dimension=384,
            metric="cosine",
            spec=ServerlessSpec(cloud="aws", region="us-east-1")
        )
    
    embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
    v_store = PineconeVectorStore(
        index_name=PINECONE_INDEX,
        embedding=embeddings,
        pinecone_api_key=PINECONE_API_KEY
    )
    language_model = ChatGroq(groq_api_key=GROQ_API_KEY, model_name="llama-3.1-8b-instant")
    
    return v_store, language_model

with st.spinner("Initializing AI Core & Vector Database..."):
    vectorstore, llm = init_backend()

# =========================
# FILE READING FUNCTIONS
# =========================
def read_pdf(file):
    return "".join([page.extract_text() or "" for page in PdfReader(file).pages])

def read_docx(file):
    return "\n".join([para.text for para in Document(file).paragraphs])

def read_csv(file):
    return pd.read_csv(file).to_string()

# =========================
# PROCESS DOCUMENTS
# =========================
if uploaded_files:
    with st.sidebar:
        st.markdown("<br>", unsafe_allow_html=True)
        process_btn = st.button("🚀 Index Documents to Pinecone", use_container_width=True)
        
    if process_btn:
        with st.status("Vectorizing Knowledge Base...", expanded=True) as status:
            st.write("Extracting text from files...")
            all_docs = []
            for file in uploaded_files:
                text = ""
                if file.name.endswith(".pdf"):
                    text = read_pdf(file)
                elif file.name.endswith(".docx"):
                    text = read_docx(file)
                elif file.name.endswith(".csv"):
                    text = read_csv(file)
                elif file.name.endswith(".txt"):
                    text = str(file.read(), "utf-8")

                if text.strip():
                    st.write(f"Chunking {file.name}...")
                    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
                    chunks = splitter.split_text(text)
                    for chunk in chunks:
                        all_docs.append(LangDocument(page_content=chunk, metadata={"source": file.name}))

            if all_docs:
                st.write("Upserting vectors to Pinecone...")
                vectorstore.add_documents(all_docs)
                status.update(label=f"Successfully indexed {len(all_docs)} chunks!", state="complete", expanded=False)
                st.sidebar.success("Indexing Complete.")

# =========================
# MAIN CHAT INTERFACE
# =========================
st.markdown("<h1 style='text-align: center;'>✨ Nexus AI</h1>", unsafe_allow_html=True)
st.markdown(f"<p style='text-align: center; opacity: 0.7; margin-top: -15px;'>Active Session: {active_user}</p>", unsafe_allow_html=True)
st.markdown("<br>", unsafe_allow_html=True)

retriever = vectorstore.as_retriever(search_type="similarity", search_kwargs={"k": 4})
qa_chain = ConversationalRetrievalChain.from_llm(
    llm=llm,
    retriever=retriever,
    memory=user_data["memory"], 
    verbose=False
)
#fixing the new bug
# Render Chat History
for role, message in user_data["chat_history"]:
    avatar = "👤" if role == "You" else "✨"
    with st.chat_message(role if role == "user" else "assistant", avatar=avatar):
        st.write(message)

# Chat Input Handler
if query := st.chat_input(f"Message Nexus AI..."):
    
    with st.chat_message("user", avatar="👤"):
        st.write(query)
    
    user_data["chat_history"].append(("You", query))
    
    with st.chat_message("assistant", avatar="✨"):
        with st.spinner("Synthesizing response..."):
            response = qa_chain.invoke({"question": query})
            answer = response["answer"]
            st.write(answer)
            
    user_data["chat_history"].append(("Bot", answer))